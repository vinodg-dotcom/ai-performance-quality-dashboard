import io
import numpy as np
import pandas as pd
import streamlit as st
import plotly.express as px

st.set_page_config(page_title="AI-Assisted Performance Quality Platform", page_icon="📈", layout="wide")

REQUIRED = ["test_name", "response_time_ms", "throughput_rps", "error_rate_pct", "concurrent_users"]

def demo_data():
    rng = np.random.default_rng(42)
    rows = []
    for run, multiplier in [("Baseline", 1.0), ("Current", 1.18)]:
        for i, endpoint in enumerate(["/login", "/search", "/checkout", "/products"]):
            rt = [240, 380, 520, 290][i] * multiplier * rng.uniform(0.92, 1.08)
            throughput = [115, 92, 58, 105][i] / (multiplier if run == "Current" else 1) * rng.uniform(0.93, 1.07)
            errors = [0.4, 0.8, 1.2, 0.3][i] + (0.7 if run == "Current" else 0) + rng.uniform(0, 0.15)
            rows.append({
                "run_name": run,
                "test_name": endpoint,
                "response_time_ms": round(rt, 1),
                "throughput_rps": round(throughput, 1),
                "error_rate_pct": round(errors, 2),
                "concurrent_users": 100
            })
    return pd.DataFrame(rows)

def validate(df):
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        return False, f"Missing required columns: {', '.join(missing)}"
    for c in REQUIRED[1:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    if df[REQUIRED[1:]].isna().any().any():
        return False, "Numeric metric columns contain blank or non-numeric values."
    if (df["response_time_ms"] < 0).any() or (df["throughput_rps"] < 0).any() or (df["error_rate_pct"] < 0).any():
        return False, "Metrics cannot be negative."
    return True, ""

def investigate(row, baseline_row=None, response_threshold=500, error_threshold=2.0, regression_threshold=10.0):
    notes = []
    if row["response_time_ms"] > response_threshold:
        notes.append(f"Response time is above the configured limit ({row['response_time_ms']:.1f} ms > {response_threshold:.1f} ms).")
    if row["error_rate_pct"] > error_threshold:
        notes.append(f"Error rate is above the configured limit ({row['error_rate_pct']:.2f}% > {error_threshold:.2f}%).")
    if baseline_row is not None:
        if baseline_row["response_time_ms"] > 0:
            delta = (row["response_time_ms"] / baseline_row["response_time_ms"] - 1) * 100
            if delta >= regression_threshold:
                notes.append(f"Response time regressed by {delta:.1f}% versus baseline.")
        if baseline_row["throughput_rps"] > 0:
            delta_t = (row["throughput_rps"] / baseline_row["throughput_rps"] - 1) * 100
            if delta_t <= -regression_threshold:
                notes.append(f"Throughput dropped by {abs(delta_t):.1f}% versus baseline.")
        if row["error_rate_pct"] - baseline_row["error_rate_pct"] >= 1:
            notes.append("Error rate increased by at least 1 percentage point versus baseline.")
    if not notes:
        notes.append("No configured threshold breach or major baseline regression detected. Review percentiles, server-side telemetry, and test validity before release.")
    notes.append("Suggested checks: correlate with CPU/memory, database waits, downstream latency, connection pools, and application logs.")
    return " ".join(notes)

st.title("📈 AI-Assisted Performance Testing & Quality Engineering")
st.caption("Portfolio MVP • Metrics dashboard • Baseline comparison • Configurable quality gates • Investigation notes")

with st.sidebar:
    st.header("Quality gate settings")
    response_limit = st.number_input("Max response time (ms)", min_value=1.0, value=500.0, step=50.0)
    error_limit = st.number_input("Max error rate (%)", min_value=0.0, value=2.0, step=0.5)
    regression_limit = st.number_input("Regression warning (%)", min_value=1.0, value=10.0, step=1.0)
    st.markdown("---")
    st.markdown("**CSV columns**")
    st.code("run_name, test_name, response_time_ms, throughput_rps, error_rate_pct, concurrent_users", language="text")

st.markdown("Upload the two JMeter CSV files separately. Raw JMeter sample-level CSVs are summarized into one row per transaction.")
base_file = st.file_uploader("Upload baseline CSV", type=["csv"], key="baseline_upload")
current_file = st.file_uploader("Upload current CSV", type=["csv"], key="current_upload")

jmeter_error_details = []

def normalize_jmeter_csv(file_obj, run_name):
    """Convert JMeter sample-level CSV into the dashboard's required schema."""
    global jmeter_error_details
    raw = pd.read_csv(file_obj)
    # Preserve individual failed JMeter samples for the dedicated errors slide.
    if {"success", "label"}.issubset(raw.columns):
        success_values = raw["success"].astype(str).str.strip().str.lower().isin(["true", "1", "yes"])
        failed_rows = raw.loc[~success_values]
        for _, failed in failed_rows.iterrows():
            response_code = str(failed.get("responseCode", "Unknown"))
            response_message = str(failed.get("responseMessage", ""))
            failure_message = str(failed.get("failureMessage", ""))
            details = f"{run_name} | {failed.get('label', 'Unknown transaction')} | response code: {response_code}"
            if response_message and response_message.lower() != "nan":
                details += f" | response: {response_message}"
            if failure_message and failure_message.lower() != "nan":
                details += f" | failure: {failure_message}"
            jmeter_error_details.append(details)
    # If already in dashboard format, retain it and assign the selected run name.
    if all(c in raw.columns for c in REQUIRED):
        normalized = raw.copy()
        normalized["run_name"] = run_name
        return normalized

    jmeter_required = {"elapsed", "label", "success", "allThreads"}
    missing = sorted(jmeter_required - set(raw.columns))
    if missing:
        raise ValueError(
            "CSV is neither in dashboard format nor a recognized JMeter sample CSV. "
            f"Missing JMeter columns: {', '.join(missing)}"
        )

    raw["elapsed"] = pd.to_numeric(raw["elapsed"], errors="coerce")
    raw["allThreads"] = pd.to_numeric(raw["allThreads"], errors="coerce")
    raw = raw.dropna(subset=["elapsed", "label"])
    if raw.empty:
        raise ValueError("No valid JMeter samples were found in this file.")

    if raw["success"].dtype != bool:
        raw["success"] = raw["success"].astype(str).str.strip().str.lower().isin(
            ["true", "1", "yes"]
        )

    # Estimate the full run span: earliest sample start to latest sample completion.
    # This is more reliable than last-start minus first-start, though JMeter's
    # listener summary can differ slightly depending on test startup/shutdown timing.
    duration_seconds = 0.0
    if "timeStamp" in raw.columns:
        raw["timeStamp"] = pd.to_numeric(raw["timeStamp"], errors="coerce")
        valid_times = raw.dropna(subset=["timeStamp"]).copy()
        if not valid_times.empty:
            run_start_ms = float(valid_times["timeStamp"].min())
            run_end_ms = float((valid_times["timeStamp"] + valid_times["elapsed"]).max())
            duration_seconds = max((run_end_ms - run_start_ms) / 1000.0, 0.0)

    if duration_seconds <= 0:
        # Fallback only when timestamps are unavailable/invalid. This is an
        # estimate, not a substitute for JMeter's official test duration.
        duration_seconds = max(float(raw["elapsed"].sum()) / 1000.0, 0.001)

    rows = []
    for label, group in raw.groupby("label", dropna=False):
        total = len(group)
        failures = int((~group["success"]).sum())
        rows.append({
            "run_name": run_name,
            "test_name": str(label),
            "sample_count": total,
            "response_time_ms": round(float(group["elapsed"].mean()), 2),
            # Transaction throughput contribution: this label's sample count / full run span.
            "throughput_rps": round(total / duration_seconds, 3),
            "error_rate_pct": round(failures / total * 100, 2),
            "concurrent_users": int(group["allThreads"].max()) if group["allThreads"].notna().any() else 1,
            "estimated_run_duration_s": round(duration_seconds, 3),
        })
    return pd.DataFrame(rows)

if base_file is not None or current_file is not None:
    if base_file is None or current_file is None:
        st.warning("Upload both baseline and current CSV files to run a comparison. You can upload either one first.")
        available = []
        if base_file is not None:
            available.append(normalize_jmeter_csv(base_file, "Baseline"))
        if current_file is not None:
            available.append(normalize_jmeter_csv(current_file, "Current"))
        data = pd.concat(available, ignore_index=True) if available else demo_data()
    else:
        try:
            baseline_data = normalize_jmeter_csv(base_file, "Baseline")
            current_data = normalize_jmeter_csv(current_file, "Current")
            data = pd.concat([baseline_data, current_data], ignore_index=True)
            valid, message = validate(data)
            if not valid:
                st.error(message)
                st.stop()
            st.success("Loaded baseline and current results. Raw JMeter samples have been summarized for comparison.")
            st.caption("Throughput is estimated from the earliest sample start to the latest sample completion. For exact agreement with JMeter's Summary Report, use the actual JMeter test duration; timestamp-derived duration can differ slightly.")
        except Exception as exc:
            st.error(f"Could not read or convert the CSV files: {exc}")
            st.stop()
else:
    data = demo_data()
    st.info("Showing sample data. Upload both JMeter CSV files to analyze your test runs.")

data["run_name"] = data["run_name"] if "run_name" in data.columns else "Current"
run_names = list(data["run_name"].dropna().unique())
if not run_names:
    st.error("No run names found.")
    st.stop()

col_a, col_b = st.columns(2)
with col_a:
    current_run = st.selectbox("Current test run", run_names, index=len(run_names)-1)
with col_b:
    baseline_options = ["None"] + [r for r in run_names if r != current_run]
    baseline_run = st.selectbox("Baseline run", baseline_options, index=1 if len(baseline_options) > 1 else 0)

current = data[data["run_name"] == current_run].copy()
baseline = data[data["run_name"] == baseline_run].copy() if baseline_run != "None" else pd.DataFrame()

# Backward compatibility for dashboard-format CSVs that lack raw-sample metadata.
if "sample_count" not in current.columns:
    current["sample_count"] = 1
if "estimated_run_duration_s" not in current.columns:
    current["estimated_run_duration_s"] = 0.0
if "sample_count" not in baseline.columns and not baseline.empty:
    baseline["sample_count"] = 1
if "estimated_run_duration_s" not in baseline.columns and not baseline.empty:
    baseline["estimated_run_duration_s"] = 0.0

avg_rt = current["response_time_ms"].mean()
avg_err = current["error_rate_pct"].mean()
users = current["concurrent_users"].max()
# Run-level throughput is total samples divided by the run's estimated duration.
# The per-transaction throughput values remain available for transaction charts.
if "sample_count" in current.columns and "estimated_run_duration_s" in current.columns:
    run_duration_s = float(current["estimated_run_duration_s"].max())
    total_samples = int(current["sample_count"].sum())
    if run_duration_s > 0:
        overall_throughput = total_samples / run_duration_s
    else:
        # Aggregated/demo-format CSVs do not carry raw timing metadata.
        overall_throughput = float(current["throughput_rps"].sum())
else:
    run_duration_s = 0.0
    total_samples = len(current)
    overall_throughput = float(current["throughput_rps"].sum())

k1, k2, k3, k4 = st.columns(4)
k1.metric("Avg response time", f"{avg_rt:.1f} ms")
k2.metric("Overall throughput (est.)", f"{overall_throughput:.2f} req/s")
k3.metric("Avg error rate", f"{avg_err:.2f}%")
k4.metric("Max concurrent users", f"{users:.0f}")

st.subheader("Performance metrics")
left, right = st.columns(2)
with left:
    st.plotly_chart(px.bar(current, x="test_name", y="response_time_ms", title="Response time by transaction (ms)", text_auto=".1f"), use_container_width=True)
with right:
    st.plotly_chart(px.bar(current, x="test_name", y="throughput_rps", title="Transaction throughput contribution (req/s)", text_auto=".1f"), use_container_width=True)

st.subheader("Baseline regression analysis")
if baseline.empty:
    st.write("Select a baseline run to compare performance changes.")
    comparison = current.copy()
    comparison["response_time_change_pct"] = np.nan
    comparison["throughput_change_pct"] = np.nan
    comparison["error_rate_change_pp"] = np.nan
else:
    base_agg = baseline.groupby("test_name", as_index=False).agg(
        baseline_response_time_ms=("response_time_ms", "mean"),
        baseline_throughput_rps=("throughput_rps", "mean"),
        baseline_error_rate_pct=("error_rate_pct", "mean")
    )
    comparison = current.merge(base_agg, on="test_name", how="left")
    comparison["response_time_change_pct"] = (comparison["response_time_ms"] / comparison["baseline_response_time_ms"] - 1) * 100
    comparison["throughput_change_pct"] = (comparison["throughput_rps"] / comparison["baseline_throughput_rps"] - 1) * 100
    comparison["error_rate_change_pp"] = comparison["error_rate_pct"] - comparison["baseline_error_rate_pct"]
    st.dataframe(comparison[["test_name", "response_time_ms", "baseline_response_time_ms", "response_time_change_pct", "throughput_rps", "baseline_throughput_rps", "throughput_change_pct", "error_rate_pct", "baseline_error_rate_pct", "error_rate_change_pp"]].round(2), use_container_width=True)

st.subheader("Quality gate result")
response_failures = int((current["response_time_ms"] > response_limit).sum())
error_failures = int((current["error_rate_pct"] > error_limit).sum())
regression_failures = 0 if baseline.empty else int((comparison["response_time_change_pct"] >= regression_limit).sum())
passed = response_failures == 0 and error_failures == 0 and regression_failures == 0
if passed:
    st.success("PASS — no configured quality-gate breaches found.")
else:
    st.error("FAIL — one or more configured quality-gate breaches found.")
g1, g2, g3 = st.columns(3)
g1.metric("Response-time breaches", response_failures)
g2.metric("Error-rate breaches", error_failures)
g3.metric("Regression warnings", regression_failures)

st.subheader("Investigation assistant")
notes = []
base_lookup = {}
if not baseline.empty:
    base_lookup = {r["test_name"]: r for _, r in baseline.groupby("test_name", as_index=False).mean(numeric_only=True).iterrows()}
for _, row in current.iterrows():
    base_row = base_lookup.get(row["test_name"])
    note = investigate(row, base_row, response_limit, error_limit, regression_limit)
    notes.append({"Transaction": row["test_name"], "Assessment": note})
for item in notes:
    with st.expander(item["Transaction"]):
        st.write(item["Assessment"])

st.subheader("Current run data")
st.dataframe(current, use_container_width=True)

st.subheader("Export and publish results")
st.caption("Generate downloadable Excel and PowerPoint reports from the data currently displayed in this dashboard.")

# Build a compact run summary for export.
export_summary = pd.DataFrame([
    {"Metric": "Selected current run", "Value": str(current_run)},
    {"Metric": "Selected baseline run", "Value": str(baseline_run)},
    {"Metric": "Transactions in current run", "Value": int(len(current))},
    {"Metric": "Average response time (ms)", "Value": round(float(current["response_time_ms"].mean()), 2)},
    {"Metric": "Overall throughput estimate (req/s)", "Value": round(float(overall_throughput), 3)},
    {"Metric": "Total samples in current run", "Value": int(total_samples)},
    {"Metric": "Estimated run duration (seconds)", "Value": round(float(run_duration_s), 3)},
    {"Metric": "Average error rate (%)", "Value": round(float(current["error_rate_pct"].mean()), 2)},
    {"Metric": "Maximum concurrent users", "Value": int(current["concurrent_users"].max())},
    {"Metric": "Response-time breaches", "Value": int(response_failures)},
    {"Metric": "Error-rate breaches", "Value": int(error_failures)},
    {"Metric": "Regression warnings", "Value": int(regression_failures)},
    {"Metric": "Quality gate", "Value": "PASS" if passed else "FAIL"},
])
export_findings = pd.DataFrame(notes) if notes else pd.DataFrame(columns=["Transaction", "Assessment"])

excel_buffer = io.BytesIO()
with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
    export_summary.to_excel(writer, sheet_name="Run Summary", index=False)
    comparison.to_excel(writer, sheet_name="Baseline Comparison", index=False)
    current.to_excel(writer, sheet_name="Current Run Data", index=False)
    if not baseline.empty:
        baseline.to_excel(writer, sheet_name="Baseline Run Data", index=False)
    export_findings.to_excel(writer, sheet_name="Investigation Notes", index=False)

# Apply basic readable formatting to the Excel workbook.
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
excel_buffer.seek(0)
workbook = load_workbook(excel_buffer)

# Add native, editable Excel charts using the baseline comparison worksheet.
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

chart_sheet = workbook.create_sheet("Performance Charts")
chart_sheet["A1"] = "Performance Test Charts"
chart_sheet["A2"] = f"Current run: {current_run} | Baseline run: {baseline_run}"
chart_sheet["A4"] = "Charts compare the current run with the selected baseline by transaction."

comparison_ws = workbook["Baseline Comparison"]
header_map = {str(cell.value): cell.column for cell in comparison_ws[1] if cell.value is not None}
chart_specs = [
    ("Response Time (ms)", "response_time_ms", "baseline_response_time_ms", 6),
    ("Throughput (req/s)", "throughput_rps", "baseline_throughput_rps", 24),
    ("Error Rate (%)", "error_rate_pct", "baseline_error_rate_pct", 42),
]
charts_added = 0
if not baseline.empty:
    for chart_title, current_col, baseline_col, top_row in chart_specs:
        if current_col not in header_map or baseline_col not in header_map:
            continue
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = chart_title
        chart.y_axis.title = chart_title
        chart.x_axis.title = "Transaction"
        chart.height = 8
        chart.width = 16
        data_ref = Reference(
            comparison_ws,
            min_col=header_map[current_col],
            max_col=header_map[current_col],
            min_row=1,
            max_row=comparison_ws.max_row,
        )
        baseline_ref = Reference(
            comparison_ws,
            min_col=header_map[baseline_col],
            max_col=header_map[baseline_col],
            min_row=1,
            max_row=comparison_ws.max_row,
        )
        cats = Reference(comparison_ws, min_col=header_map["test_name"], min_row=2, max_row=comparison_ws.max_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.add_data(baseline_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.legend.position = "b"
        chart_sheet.add_chart(chart, f"A{top_row}")
        charts_added += 1

if charts_added == 0:
    chart_sheet["A6"] = "Charts require both a baseline and current run. Upload/select both runs to generate comparison charts."

for worksheet in workbook.worksheets:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[1].height = 28
    for column_cells in worksheet.columns:
        width = min(max(max(len(str(c.value or "")) for c in column_cells) + 2, 12), 42)
        worksheet.column_dimensions[column_cells[0].column_letter].width = width
    for row_cells in worksheet.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
if "Performance Charts" in workbook.sheetnames:
    workbook.active = workbook.sheetnames.index("Performance Charts")
    workbook["Performance Charts"].sheet_view.showGridLines = False
    workbook["Performance Charts"].sheet_properties.pageSetUpPr.fitToPage = True

formatted_excel = io.BytesIO()
workbook.save(formatted_excel)
formatted_excel.seek(0)

# Build a PowerPoint report in memory so the dashboard can offer a download.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ppt_buffer = io.BytesIO()
presentation = Presentation()
presentation.slide_width = Inches(13.333)
presentation.slide_height = Inches(7.5)

def add_export_slide_title(slide, title, subtitle=""):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.65))
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = title
    title_paragraph.font.size = Pt(26)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = RGBColor(23, 54, 93)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.62), Inches(1.05), Inches(12.0), Inches(0.35))
        sub_paragraph = sub_box.text_frame.paragraphs[0]
        sub_paragraph.text = subtitle
        sub_paragraph.font.size = Pt(11)
        sub_paragraph.font.color.rgb = RGBColor(90, 90, 90)

def add_export_bullets(slide, items, top=1.65, font_size=18):
    """Create readable bullets with an explicit gap after the bullet glyph."""
    box = slide.shapes.add_textbox(Inches(0.78), Inches(top), Inches(11.9), Inches(5.3))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.10)
    frame.margin_right = Inches(0.10)
    frame.margin_top = Inches(0.08)
    frame.margin_bottom = Inches(0.08)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        # Use a literal bullet plus spaces so the gap renders consistently in PowerPoint.
        paragraph.text = f"•   {str(item)}"
        paragraph.font.name = 'Aptos'
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = RGBColor(55, 60, 68)
        paragraph.space_before = Pt(2)
        paragraph.space_after = Pt(15)
        paragraph.line_spacing = 1.12

# Create chart images for the presentation, based on the selected current and baseline runs.
def add_chart_slide(title, subtitle, chart_kind):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(10.8, 4.4))
    labels = current["test_name"].astype(str).tolist()
    x = np.arange(len(labels))
    width = 0.36 if not baseline.empty else 0.62
    if chart_kind == "response":
        current_values = [float(current.loc[current["test_name"].astype(str) == label, "response_time_ms"].mean()) for label in labels]
        ax.bar(x + (width / 2 if not baseline.empty else 0), current_values, width=width, label="Current average response time (ms)")
        if not baseline.empty:
            b = baseline.groupby("test_name")["response_time_ms"].mean()
            baseline_values = [float(b.get(label, np.nan)) for label in labels]
            ax.bar(x - width / 2, baseline_values, width=width, label="Baseline average response time (ms)")
        ax.set_ylabel("Average response time (ms)")
    else:
        current_values = [float(current.loc[current["test_name"].astype(str) == label, "throughput_rps"].mean()) for label in labels]
        ax.bar(x + (width / 2 if not baseline.empty else 0), current_values, width=width, label="Current transaction throughput contribution (req/s)")
        if not baseline.empty:
            b = baseline.groupby("test_name")["throughput_rps"].mean()
            baseline_values = [float(b.get(label, np.nan)) for label in labels]
            ax.bar(x - width / 2, baseline_values, width=width, label="Baseline transaction throughput contribution (req/s)")
        ax.set_ylabel("Transaction throughput contribution (req/s)")
    ax.set_xticks(x)
    ax.set_xticklabels(labels, rotation=18, ha="right")
    ax.set_title(title)
    ax.grid(axis="y", alpha=0.25)
    ax.set_axisbelow(True)
    ax.legend(loc="best", fontsize=9)
    fig.tight_layout()
    img = io.BytesIO()
    fig.savefig(img, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    img.seek(0)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_export_slide_title(slide, title, subtitle)
    slide.shapes.add_picture(img, Inches(0.8), Inches(1.45), width=Inches(11.8), height=Inches(5.4))


slide = presentation.slides.add_slide(presentation.slide_layouts[6])
add_export_slide_title(slide, "Performance Test Results", "AI-Assisted Performance Testing & Quality Engineering")
add_export_bullets(slide, [
    f"Current run: {current_run} ({total_samples} samples)",
    f"Baseline run: {baseline_run}",
    f"Transactions in current run: {current['test_name'].nunique()}",
    f"Average response time across transaction averages: {current['response_time_ms'].mean():.2f} ms",
    f"Overall throughput estimate: {overall_throughput:.3f} req/s",
    f"Estimated run duration: {run_duration_s:.2f} seconds",
    f"Average transaction error rate: {current['error_rate_pct'].mean():.2f}%",
    f"Quality gate: {'PASS' if passed else 'FAIL'}",
])


add_chart_slide("Response Time Comparison", "Grouped bars compare average response time by transaction", "response")
add_chart_slide("Transaction Throughput Contribution", "Grouped bars show each transaction sample count divided by the estimated full run duration", "throughput")

slide = presentation.slides.add_slide(presentation.slide_layouts[6])
add_export_slide_title(slide, "Baseline Comparison", "Selected metrics; full comparison remains available in the Excel export")
if baseline.empty:
    add_export_bullets(slide, ["No baseline was selected; percentage changes are not available."])
else:
    # Keep the slide table readable; the Excel export retains the full comparison columns.
    display_cols = [col for col in [
        "test_name", "response_time_ms", "baseline_response_time_ms",
        "response_time_change_pct", "throughput_rps", "baseline_throughput_rps",
        "error_rate_pct"
    ] if col in comparison.columns]
    table_df = comparison[display_cols].head(10)
    rows_count, cols_count = len(table_df) + 1, len(display_cols)
    table = slide.shapes.add_table(rows_count, cols_count, Inches(0.45), Inches(1.55), Inches(12.4), Inches(4.8)).table
    for j, col in enumerate(display_cols):
        table.cell(0, j).text = col.replace("_", " ")
    for i, (_, row) in enumerate(table_df.iterrows(), start=1):
        for j, col in enumerate(display_cols):
            value = row[col]
            table.cell(i, j).text = "" if pd.isna(value) else (f"{value:.2f}" if isinstance(value, (int, float, np.number)) else str(value))
    for cell in table.rows[0].cells:
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(9)
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(23, 54, 93)
    for row_index in range(1, len(table.rows)):
        row = table.rows[row_index]
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(8)

# Dedicated errors slide: list every failed JMeter sample when available; otherwise state clearly that none were recorded.
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
add_export_slide_title(slide, "Errors & Failed Transactions", "Failed samples from uploaded JMeter results and transaction-level error rates")
error_items = []
if jmeter_error_details:
    error_items.append(f"Total failed JMeter samples recorded across uploaded files: {len(jmeter_error_details)}")
    error_items.extend(jmeter_error_details)
else:
    current_error_transactions = current[current["error_rate_pct"] > 0]
    if not current_error_transactions.empty:
        error_items.append(f"Transactions with non-zero error rate: {len(current_error_transactions)}")
        for _, err_row in current_error_transactions.iterrows():
            error_items.append(f"{err_row['test_name']}: error rate {err_row['error_rate_pct']:.2f}%")
        error_items.append("Detailed response codes/messages are not present in the uploaded normalized CSV.")
    else:
        error_items.append("No errors recorded in the uploaded current-run results.")
        error_items.append("Current-run transaction error rates are all 0.00%.")
# Split long error lists over multiple slides so every recorded error is included and readable.
error_chunks = [error_items[i:i+12] for i in range(0, len(error_items), 12)] or [["No errors recorded in the uploaded current-run results."]]
for chunk_index, chunk in enumerate(error_chunks):
    if chunk_index == 0:
        error_slide = slide
    else:
        error_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_export_slide_title(error_slide, f"Errors & Failed Transactions (continued {chunk_index})", "All recorded failed samples are listed; long reports continue on additional slides")
    add_export_bullets(error_slide, chunk, top=1.55, font_size=14)

slide = presentation.slides.add_slide(presentation.slide_layouts[6])
add_export_slide_title(slide, "Quality Gate & Investigation Notes", "Rule-based triage guidance; not an automated root-cause diagnosis")
findings_items = [
    f"Overall quality gate: {'PASS' if passed else 'FAIL'}",
    f"Response-time threshold: {response_limit:.1f} ms; transactions breaching it: {response_failures}",
    f"Error-rate threshold: {error_limit:.2f}%; transactions breaching it: {error_failures}",
    f"Baseline regression threshold: {regression_limit:.1f}%; transactions flagged: {regression_failures}",
]
# Explicitly name each transaction breaching configured thresholds or showing a baseline regression.
failed_transactions = []
for _, r in current.iterrows():
    reasons = []
    if r["response_time_ms"] > response_limit:
        reasons.append(f"response time {r['response_time_ms']:.2f} ms exceeds {response_limit:.1f} ms")
    if r["error_rate_pct"] > error_limit:
        reasons.append(f"error rate {r['error_rate_pct']:.2f}% exceeds {error_limit:.2f}%")
    match = comparison[comparison["test_name"] == r["test_name"]]
    if not baseline.empty and not match.empty and pd.notna(match.iloc[0].get("response_time_change_pct")) and match.iloc[0]["response_time_change_pct"] >= regression_limit:
        reasons.append(f"response time regressed {match.iloc[0]['response_time_change_pct']:.1f}% vs baseline")
    if reasons:
        failed_transactions.append(f"{r['test_name']}: " + "; ".join(reasons))
if failed_transactions:
    findings_items.append("Flagged transactions:")
    findings_items.extend(failed_transactions[:6])
else:
    findings_items.append("No transactions breached the configured response-time or error-rate thresholds, and no response-time regression exceeded the warning limit.")
findings_items.append("Recommended follow-up: validate the load profile and repeat the run; correlate any flagged transaction with CPU, memory, database waits, downstream latency, connection pools, and application logs.")
add_export_bullets(slide, findings_items, font_size=13)
presentation.save(ppt_buffer)
ppt_buffer.seek(0)

download_col1, download_col2 = st.columns(2)
with download_col1:
    st.download_button(
        "Download Excel report (.xlsx)",
        data=formatted_excel.getvalue(),
        file_name="performance_test_report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
with download_col2:
    st.download_button(
        "Download PowerPoint report (.pptx)",
        data=ppt_buffer.getvalue(),
        file_name="performance_test_results.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        use_container_width=True,
    )

st.markdown("---")
st.caption("Important: this MVP uses rule-based investigation notes, not a diagnostic AI model. It supports triage and does not prove root cause.")
